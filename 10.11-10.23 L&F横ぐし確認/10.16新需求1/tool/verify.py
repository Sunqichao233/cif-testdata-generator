# -*- coding: utf-8 -*-
"""
生成した対比表の突合列（EXACT / TRIM / AND）を実際に評価して、
どのセルが TRUE / FALSE になるかを先に確かめる。
"""
import re, sys
import openpyxl

path = sys.argv[1] if len(sys.argv) > 1 else 'test2.xlsx'
wb = openpyxl.load_workbook(path)

REF = re.compile(r'^\$?([A-Z]{1,3})\$?(\d+)$')
RNG = re.compile(r'^\$?([A-Z]{1,3})\$?(\d+):\$?([A-Z]{1,3})\$?(\d+)$')
from openpyxl.utils import column_index_from_string as ci, get_column_letter as CL


def tokenize(s):
    out, i = [], 0
    while i < len(s):
        c = s[i]
        if c == '"':
            j = i + 1
            buf = ''
            while j < len(s):
                if s[j] == '"':
                    if j + 1 < len(s) and s[j + 1] == '"':
                        buf += '"'; j += 2; continue
                    break
                buf += s[j]; j += 1
            out.append(('str', buf)); i = j + 1
        elif c in '&(),':
            out.append((c, c)); i += 1
        elif c == ' ':
            i += 1
        else:
            j = i
            while j < len(s) and s[j] not in '&(),"':
                j += 1
            out.append(('word', s[i:j].strip())); i = j
    return out


class P:
    def __init__(self, toks, ws, depth):
        self.t, self.i, self.ws, self.d = toks, 0, ws, depth

    def peek(self):
        return self.t[self.i] if self.i < len(self.t) else (None, None)

    def take(self):
        v = self.t[self.i]; self.i += 1; return v

    def expr(self):
        v = self.atom()
        while self.peek()[0] == '&':
            self.take()
            v = str(v) + str(self.atom())
        return v

    def atom(self):
        k, v = self.take()
        if k == 'str':
            return v
        if k == '(':
            r = self.expr(); self.take(); return r
        if k == 'word':
            if self.peek()[0] == '(':
                return self.call(v.upper())
            return cell(self.ws, v, self.d)
        raise ValueError('token %r' % ((k, v),))

    def call(self, name):
        self.take()                      # (
        args = []
        if self.peek()[0] != ')':
            args.append(self.expr())
            while self.peek()[0] == ',':
                self.take(); args.append(self.expr())
        self.take()                      # )
        if name == 'SUBSTITUTE':
            return str(args[0]).replace(str(args[1]), str(args[2]))
        if name == 'TRIM':
            return re.sub(r' +', ' ', str(args[0])).strip()
        if name == 'EXACT':
            return str(args[0]) == str(args[1])
        if name == 'AND':
            return all(bool(x) if isinstance(x, bool) else str(x).upper() == 'TRUE'
                       for x in args[0])
        if name == 'IF':
            return args[1] if args[0] else args[2]
        raise ValueError('未対応の関数 ' + name)


def cell(ws, ref, depth):
    m = RNG.match(ref)
    if m:
        c1, r1, c2, r2 = ci(m.group(1)), int(m.group(2)), ci(m.group(3)), int(m.group(4))
        return [val(ws, '%s%d' % (CL(c), r1), depth)
                for c in range(c1, c2 + 1) for r in (r1,)]
    if REF.match(ref):
        return val(ws, ref, depth)
    if ref.upper() in ('TRUE', 'FALSE'):
        return ref.upper() == 'TRUE'
    return ref


def val(ws, a, depth=0):
    v = ws[a.replace('$', '')].value
    if v is None:
        return ''
    if isinstance(v, str) and v.startswith('=') and depth < 10:
        return ev(ws, v, depth + 1)
    if isinstance(v, bool):
        return v
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        return '%g' % v
    return str(v)


def ev(ws, f, depth=0):
    return P(tokenize(f[1:]), ws, depth).expr()


tot = {}
bad = []
for n in wb.sheetnames:
    ws = wb[n]
    hdr = {}
    for c in range(1, ws.max_column + 1):
        v = str(ws.cell(row=3, column=c).value or '')
        if v:
            hdr.setdefault(v, c)
    cAll, cSql, cWhy = hdr.get('全項目一致'), hdr.get('SQL文一致'), hdr.get('相違した項目')
    for key, col in (('①②（項目）', cAll), ('①③（SQL文）', cSql)):
        if not col:
            continue
        for r in range(4, ws.max_row + 1):
            f = ws.cell(row=r, column=col).value
            if not (isinstance(f, str) and f.startswith('=')):
                continue
            res = ev(ws, f)
            res = bool(res) if isinstance(res, bool) else str(res).upper() == 'TRUE'
            tot.setdefault(key, [0, 0])
            tot[key][0 if res else 1] += 1
            if not res:
                why = ''
                if cWhy and key.startswith('①②'):
                    why = str(ev(ws, ws.cell(row=r, column=cWhy).value) or '')
                bad.append((n, r, key, why))

print('=== 評価結果 ===')
for k, (t, f) in tot.items():
    print('  %-12s TRUE %5d / FALSE %4d' % (k, t, f))
print()
seen = {}
for n, r, k, why in bad:
    seen.setdefault((n, k, why), []).append(r)
print('FALSE になった行')
for (n, k, why), rs in seen.items():
    print('  %-26s %-12s %3d件  行%-26s %s' % (n, k, len(rs), str(rs[:6]), why))
