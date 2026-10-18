# -*- coding: utf-8 -*-
"""
insert.sql を「-- N_xxxマスタ」のコメント単位でシートに分割し、
SQL文を項目ごとに切り分けた Excel を作る。
様式は サンプル.xlsx（シート insert）に合わせている。

  A列        原文SQL（1行そのまま）
  区切りBLK  カンマで単純分割（サンプルと同じ切り方）
  値BLK      引用符・to_timestamp を外した素の値  ← マスタ定義書と対比する欄
  再構築BLK  構文＋テーブル名から原文を組み立て直し、EXACT で一致を検証

使い方 : python sql2excel.py [insert.sql] [出力.xlsx]
"""
import os, re, sys, datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL

# ---------------- 設定 ----------------
FONT      = '游ゴシック'
SZ        = 9
C_RAW     = 'FFD9D9D9'   # 原文ブロック
C_SPLIT   = 'FFDDEBF7'   # カンマ区切りブロック
C_VAL     = 'FFFFFF99'   # 値ブロック（サンプルと同じ色）
C_REBUILD = 'FFE2EFDA'   # 再構築・検証ブロック
C_NOTE    = 'FFFCE4D6'   # 判定・備考
C_OTHER   = 'FFFFF2CC'   # INSERT以外の文（DELETE / UPDATE）
SEPW      = 2.0          # 区切り列の幅

thin = Side(style='thin', color='FFBFBFBF')
BOX  = Border(left=thin, right=thin, top=thin, bottom=thin)

Q = chr(39)   # シングルクォート


# ---------------- SQL の解析 ----------------
def split_args(body):
    """カンマで分割。'...' の中と ( ) の中のカンマは区切りにしない。"""
    out, buf, q, d, i = [], '', False, 0, 0
    while i < len(body):
        ch = body[i]
        if q:
            if ch == Q:
                if i + 1 < len(body) and body[i + 1] == Q:
                    buf += Q + Q
                    i += 2
                    continue
                q = False
            buf += ch
        else:
            if ch == Q:
                q = True
                buf += ch
            elif ch == '(':
                d += 1
                buf += ch
            elif ch == ')':
                d -= 1
                buf += ch
            elif ch == ',' and d == 0:
                out.append(buf)
                buf = ''
            else:
                buf += ch
        i += 1
    out.append(buf)
    return [x.strip() for x in out]


RE_TS = re.compile(r"^to_timestamp\(\s*'(.*?)'\s*,\s*'(.*?)'\s*\)$", re.I)


def to_value(tok):
    """SQL のトークンを Excel に入れる値にする。戻り値 (値, 種別)"""
    t = tok.strip()
    if t == '':
        return ('', '空')
    if t.upper() == 'NULL':
        return ('NULL', 'NULL')
    m = RE_TS.match(t)
    if m:
        return (m.group(1), '日時')          # 書式文字列は捨てて日時だけ残す
    if t.startswith(Q) and t.endswith(Q) and len(t) >= 2:
        return (t[1:-1].replace(Q + Q, Q), '文字列')
    if re.match(r'^-?\d+$', t):
        return (int(t), '数値')
    if re.match(r'^-?\d+\.\d+$', t):
        return (float(t), '数値')
    return (t, 'その他')


RE_INS_VAL = re.compile(r'^(INSERT\s+INTO\s+)(\S+)(\s+VALUES\s*\()(.*)(\)\s*;)\s*$', re.I)
RE_INS_SEL = re.compile(
    r"^(INSERT\s+INTO\s+)(\S+)(\s+SELECT\s+)(\S+?)\s*,\s*(\S+)"
    r"(\s+FROM\s+)(\S+)(\s+WHERE\s+)(\S+)(\s*=\s*)'(.*?)'(\s*;)\s*$", re.I)


def parse(line):
    """1行を種別ごとに分解する。"""
    t = line.strip()
    m = RE_INS_VAL.match(t)
    if m:
        pre, tbl, mid, body, suf = m.groups()
        return dict(kind='VALUES', raw=t, tbl=tbl, pre=pre, mid=mid, body=body, suf=suf,
                    naive=body.split(','), args=split_args(body))
    m = RE_INS_SEL.match(t)
    if m:
        (pre, tbl, s_sel, col1, col2, s_from, src, s_where,
         wcol, s_eq, wval, suf) = m.groups()
        return dict(kind='SELECT', raw=t, tbl=tbl,
                    pre=pre, s_sel=s_sel, col1=col1, col2=col2, s_from=s_from,
                    src=src, s_where=s_where, wcol=wcol, s_eq=s_eq, wval=wval, suf=suf)
    k = t.split()[0].upper().rstrip(';') if t.split() else ''
    return dict(kind=k or 'その他', raw=t, tbl='')


def load(path):
    lines = open(path, encoding='utf-8').read().splitlines()
    secs, cur, head = [], None, []
    for no, line in enumerate(lines, 1):
        t = line.strip()
        if not t:
            continue
        m = re.match(r'^--\s*(\S.*?)\s*$', t)
        if m:
            cur = dict(name=m.group(1), line=no, stmts=[])
            secs.append(cur)
            continue
        if cur is None:
            head.append((no, t))            # BEGIN; など、最初のコメントより前
            continue
        st = parse(t)
        st['line'] = no
        cur['stmts'].append(st)
    return head, secs


# ---------------- Excel の組み立て ----------------
def sheet_name(name, used):
    s = re.sub(r'[\[\]:*?/\\]', '_', name)[:31]
    base, i = s, 2
    while s in used:
        tail = '_%d' % i
        s = base[:31 - len(tail)] + tail
        i += 1
    used.add(s)
    return s


def put(ws, r, c, v, fill=None, bold=False, wrap=False, align=None, fmt=None):
    cell = ws.cell(row=r, column=c)
    cell.value = v
    cell.font = Font(name=FONT, size=SZ, bold=bold)
    cell.border = BOX
    if fill:
        cell.fill = PatternFill('solid', fgColor=fill)
    cell.alignment = Alignment(horizontal=align, vertical='center', wrap_text=wrap)
    if fmt:
        cell.number_format = fmt
    elif isinstance(v, str) and not v.startswith('='):
        cell.number_format = '@'       # '0009' の先頭ゼロを守る
    return cell


def block_label(ws, r, c1, c2, text, fill):
    for c in range(c1, c2 + 1):
        ws.cell(row=r, column=c).fill = PatternFill('solid', fgColor=fill)
        ws.cell(row=r, column=c).border = BOX
    if c2 > c1:
        ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
    put(ws, r, c1, text, fill=fill, bold=True, align='center')


def build_sheet(wb, sec, unit_names):
    ws = wb.create_sheet(sec['sheet'])
    sts = sec['stmts']
    vals = [s for s in sts if s['kind'] == 'VALUES']
    sels = [s for s in sts if s['kind'] == 'SELECT']
    is_sel = len(sels) > len(vals)

    HDR = 2          # 1行目=ブロック見出し / 2行目=列見出し
    TOP = HDR + 1    # データ開始行

    # ----- 列の割り当て -----
    if is_sel:
        split_hd = ['挿入先テーブル', '挿入列1', '挿入列2の値', '参照テーブル', '条件列', '条件値']
        val_hd   = ['部店コード', '部店名（5_部店マスタより）', 'フィルタリング設定ID']
    else:
        n_split = max([len(s['naive']) for s in vals], default=1)
        n_val   = max([len(s['args'])  for s in vals], default=1)
        split_hd = ['区切%d' % (i + 1) for i in range(n_split)]
        val_hd   = ['列%d' % (i + 1) for i in range(n_val)]

    cA   = 1
    cS0  = 3                      # 区切りブロック先頭（B列は区切り）
    cV0  = cS0 + len(split_hd) + 1
    cR0  = cV0 + len(val_hd) + 1  # 再構築ブロック先頭
    reb_hd = ['構文1', 'テーブル名', '構文2', '構文3', '構文4', '再構築SQL', '一致']
    cN0  = cR0 + len(reb_hd)      # 判定・備考
    note_hd = ['判定', '備考']
    cEnd = cN0 + len(note_hd) - 1

    # ----- 見出し -----
    block_label(ws, 1, cA, cA, '原文', C_RAW)
    block_label(ws, 1, cS0, cS0 + len(split_hd) - 1,
                '構文の切り分け' if is_sel else 'カンマ区切り（サンプルと同じ切り方）', C_SPLIT)
    block_label(ws, 1, cV0, cV0 + len(val_hd) - 1, '値　← マスタ定義書と対比する欄', C_VAL)
    block_label(ws, 1, cR0, cR0 + len(reb_hd) - 1, '再構築・検証', C_REBUILD)
    block_label(ws, 1, cN0, cEnd, '確認結果', C_NOTE)

    put(ws, HDR, cA, '原文SQL', fill=C_RAW, bold=True, align='center')
    ws.cell(row=HDR, column=2).border = BOX
    for i, h in enumerate(split_hd):
        put(ws, HDR, cS0 + i, h, fill=C_SPLIT, bold=True, align='center', wrap=True)
    for i, h in enumerate(val_hd):
        put(ws, HDR, cV0 + i, h, fill=C_VAL, bold=True, align='center', wrap=True)
    for i, h in enumerate(reb_hd):
        put(ws, HDR, cR0 + i, h, fill=C_REBUILD, bold=True, align='center')
    for i, h in enumerate(note_hd):
        put(ws, HDR, cN0 + i, h, fill=C_NOTE, bold=True, align='center')

    # ----- データ -----
    r = TOP
    for s in sts:
        put(ws, r, cA, s['raw'])

        if s['kind'] == 'VALUES':
            for i, tok in enumerate(s['naive']):
                put(ws, r, cS0 + i, tok)
            for i, tok in enumerate(s['args']):
                put(ws, r, cV0 + i, to_value(tok)[0])
            put(ws, r, cR0 + 0, s['pre'])
            put(ws, r, cR0 + 1, s['tbl'])
            put(ws, r, cR0 + 2, s['mid'])
            put(ws, r, cR0 + 3, '')
            put(ws, r, cR0 + 4, s['suf'])
            # 構文1 & テーブル名 & 構文2 & 区切り列をカンマで連結 & 構文3 & 構文4
            join = ('&"' + ',' + '"&').join('%s%d' % (CL(cS0 + i), r)
                                            for i in range(len(s['naive'])))
            put(ws, r, cR0 + 5, '={0}{5}&{1}{5}&{2}{5}&{6}&{3}{5}&{4}{5}'.format(
                CL(cR0), CL(cR0 + 1), CL(cR0 + 2), CL(cR0 + 3), CL(cR0 + 4), r, join))
            put(ws, r, cR0 + 6,
                '=IF(EXACT(%s%d,%s%d),"OK","差異")' % (CL(cA), r, CL(cR0 + 5), r),
                align='center')

        elif s['kind'] == 'SELECT':
            for i, v in enumerate([s['tbl'], s['col1'], s['col2'], s['src'], s['wcol'], s['wval']]):
                put(ws, r, cS0 + i, v)
            code = s['wval']
            put(ws, r, cV0 + 0, code)
            put(ws, r, cV0 + 1, unit_names.get(code, '（5_部店マスタに無し）'))
            put(ws, r, cV0 + 2, to_value(s['col2'])[0])
            put(ws, r, cR0 + 0, s['pre'])
            put(ws, r, cR0 + 1, s['tbl'])
            put(ws, r, cR0 + 2, s['s_sel'] + s['col1'] + ', ')
            put(ws, r, cR0 + 3, s['s_from'] + s['src'] + s['s_where'] + s['wcol'] + s['s_eq'])
            put(ws, r, cR0 + 4, s['suf'])
            # 構文1&テーブル名&構文2&ID&構文3&'部店コード'&構文4
            put(ws, r, cR0 + 5,
                '={0}{6}&{1}{6}&{2}{6}&{3}{6}&{4}{6}&"{8}"&{7}{6}&"{8}"&{5}{6}'.format(
                    CL(cR0), CL(cR0 + 1), CL(cR0 + 2), CL(cV0 + 2), CL(cR0 + 3),
                    CL(cR0 + 4), r, CL(cV0), Q))
            put(ws, r, cR0 + 6,
                '=IF(EXACT(%s%d,%s%d),"OK","差異")' % (CL(cA), r, CL(cR0 + 5), r),
                align='center')

        else:   # DELETE / UPDATE など
            for c in range(cA, cEnd + 1):
                ws.cell(row=r, column=c).fill = PatternFill('solid', fgColor=C_OTHER)
                ws.cell(row=r, column=c).border = BOX
            put(ws, r, cA, s['raw'], fill=C_OTHER)
            put(ws, r, cN0 + 1, '※%s文（INSERT以外）' % s['kind'], fill=C_OTHER)

        ws.cell(row=r, column=2).border = BOX
        r += 1

    last = r - 1

    # ----- 見ための調整 -----
    ws.freeze_panes = '%s%d' % (CL(cS0), TOP)
    ws.auto_filter.ref = '%s%d:%s%d' % (CL(cA), HDR, CL(cEnd), last)
    ws.column_dimensions['A'].width = 78
    ws.column_dimensions['B'].width = SEPW
    ws.column_dimensions[CL(cV0 - 1)].width = SEPW
    ws.column_dimensions[CL(cR0 - 1)].width = SEPW
    for i in range(len(split_hd)):
        ws.column_dimensions[CL(cS0 + i)].width = 24
    for i in range(len(val_hd)):
        ws.column_dimensions[CL(cV0 + i)].width = 26 if is_sel else 15
    for i, w in enumerate([13, 24, 11, 22, 7, 78, 6]):
        ws.column_dimensions[CL(cR0 + i)].width = w
    for i, w in enumerate([8, 34]):
        ws.column_dimensions[CL(cN0 + i)].width = w
    ws.row_dimensions[1].height = 20
    ws.row_dimensions[HDR].height = 28
    ws.sheet_view.zoomScale = 90

    sec['rows'] = (TOP, last)
    sec['ncol'] = len(val_hd)
    return ws


def build_index(wb, head, secs, src, note_rows):
    ws = wb.create_sheet('00_目次', 0)
    c = put(ws, 1, 1, 'insert.sql ⇔ マスタ定義書　対比表', bold=True)
    c.font = Font(name=FONT, size=14, bold=True)
    c.border = Border()
    rows = [
        ('対象ファイル', src),
        ('作成日', datetime.date.today().strftime('%Y/%m/%d')),
        ('シート構成', 'insert.sql のコメント行「-- N_xxx」単位で1シート。並びはファイルの記載順。'),
        ('使い方', '黄色の「値」ブロックの列見出し（列1・列2…）を、マスタ定義書の列名に書き換えて使う。'
                   'その下に入っている値を定義書・実機DBと突き合わせ、右端の「判定」「備考」に記入する。'),
        ('検証列', '「再構築SQL」は切り分けた項目から原文を組み立て直したもの。'
                   '「一致」が全て OK なら、切り分けで情報が落ちていないことの裏付けになる。'),
        ('先頭/末尾', ' / '.join('%d行: %s' % (n, t) for n, t in head)),
    ]
    r = 3
    for k, v in rows:
        put(ws, r, 1, k, fill=C_RAW, bold=True)
        for cc in range(2, 10):
            ws.cell(row=r, column=cc).border = BOX
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=9)
        put(ws, r, 2, v, wrap=True)
        ws.row_dimensions[r].height = 30
        r += 1

    r += 1
    hdr = ['No', 'シート名', 'セクション（SQLのコメント）', 'テーブル名',
           'INSERT件数', '項目数', 'SQL行', '形式', '備考']
    for i, h in enumerate(hdr, 1):
        put(ws, r, i, h, fill=C_SPLIT, bold=True, align='center', wrap=True)
    hr = r
    r += 1
    total = 0
    for i, s in enumerate(secs, 1):
        n = len([x for x in s['stmts'] if x['kind'] in ('VALUES', 'SELECT')])
        o = [x for x in s['stmts'] if x['kind'] not in ('VALUES', 'SELECT')]
        total += n
        put(ws, r, 1, i, align='center')
        c = put(ws, r, 2, s['sheet'])
        c.hyperlink = "#'%s'!A1" % s['sheet']
        c.font = Font(name=FONT, size=SZ, color='FF0563C1', underline='single')
        put(ws, r, 3, s['name'])
        put(ws, r, 4, ','.join(sorted({x['tbl'] for x in s['stmts'] if x['tbl']})))
        put(ws, r, 5, n, align='center')
        put(ws, r, 6, s['ncol'], align='center')
        put(ws, r, 7, '%d〜%d' % (s['line'], s['stmts'][-1]['line'] if s['stmts'] else s['line']),
            align='center')
        put(ws, r, 8, 'INSERT-SELECT' if any(x['kind'] == 'SELECT' for x in s['stmts'])
            else 'INSERT-VALUES', align='center')
        put(ws, r, 9, '／'.join('%s文 %d件' % (k, sum(1 for x in o if x['kind'] == k))
                                for k in sorted({x['kind'] for x in o})))
        r += 1
    put(ws, r, 2, '合計', fill=C_RAW, bold=True)
    put(ws, r, 3, '%d シート' % len(secs), fill=C_RAW)
    put(ws, r, 5, total, fill=C_RAW, bold=True, align='center')
    for cc in (1, 4, 6, 7, 8, 9):
        put(ws, r, cc, '', fill=C_RAW)
    ws.auto_filter.ref = 'A%d:I%d' % (hr, r - 1)

    r += 2
    block_label(ws, r, 1, 9, '気になる点（SQLを読んだ時点で見つかったもの）', C_NOTE)
    r += 1
    put(ws, r, 1, 'No', fill=C_SPLIT, bold=True, align='center')
    put(ws, r, 2, 'SQL行', fill=C_SPLIT, bold=True, align='center')
    for cc in range(3, 10):
        put(ws, r, cc, '', fill=C_SPLIT)
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
    put(ws, r, 3, '内容', fill=C_SPLIT, bold=True, align='center')
    ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=9)
    put(ws, r, 6, '確認したいこと', fill=C_SPLIT, bold=True, align='center')
    r += 1
    for i, (ln, what, ask) in enumerate(note_rows, 1):
        for cc in range(1, 10):
            ws.cell(row=r, column=cc).border = BOX
        put(ws, r, 1, i, align='center')
        put(ws, r, 2, str(ln), align='center')
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
        put(ws, r, 3, what, wrap=True)
        ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=9)
        put(ws, r, 6, ask, wrap=True)
        ws.row_dimensions[r].height = 34
        r += 1

    for col, w in zip('ABCDEFGHI', [6, 30, 30, 28, 11, 8, 12, 15, 30]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A%d' % (hr + 1)
    ws.sheet_view.showGridLines = False
    return ws


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else 'insert.sql'
    out = sys.argv[2] if len(sys.argv) > 2 else 'insert_定義書対比.xlsx'
    head, secs = load(src)

    # 部店コード → 部店名（5_部店マスタ の MST_UNIT_TBL から拾う）
    unit_names = {}
    for s in secs:
        for st in s['stmts']:
            if st['kind'] == 'VALUES' and st['tbl'].upper() == 'MST_UNIT_TBL':
                a = [to_value(x)[0] for x in st['args']]
                if len(a) >= 3:
                    unit_names[str(a[1])] = str(a[2])

    note_rows = [
        (317, "to_timestamp('' , 'yyyy/MM/dd hh24:mi:ss') と、日時が空文字になっている"
              "（LST_REG_TBL の最終項目）",
         'NULL にすべきところではないか。空文字のままで投入してエラーにならないか。'),
        (317, "同じ行の1つ目だけ書式が小文字 'yyyy/mm/dd hh24:mi:ss'（他の63箇所は 'yyyy/MM/dd'）",
         'PostgreSQL では同じ意味だが、記述を揃えなくてよいか。'),
        ('4〜6', "プロファイル名が 'L＆F管理者…'（全角＆）と 'L&F担当者'（半角&）で混在している",
         'ヒアリングシート・設計書側の表記はどちらか。画面に出る文字列なので要確認。'),
        ('320〜619', '25部店すべてが同じ12個の設定ID（1,4,5,7,8,9,11,12,13,14,15,16）。'
                     '2・3・6・10 だけ設定されていない',
         '抜けている4つは意図的な除外か。ヒアリングシートで聞いた設定内容と一致しているか要確認。'),
        ('56〜57', 'DELETE で PRF_ID=4・5 から EXEC_AUTH_ID 9,20,21,22,26,30 の権限を外している',
         '既存プロファイルの権限を落とす変更。ヒアリングで聞いた権限と整合しているか要確認。'),
    ]

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    used = set()
    for s in secs:
        s['sheet'] = sheet_name(s['name'], used)
    for s in secs:
        build_sheet(wb, s, unit_names)
    build_index(wb, head, secs, os.path.basename(src), note_rows)
    wb.active = 0
    wb.save(out)

    print('出力 :', out)
    print('シート数 :', len(wb.sheetnames), '（00_目次 を含む）')
    print()
    print('%-30s %-28s %5s %5s' % ('シート', 'テーブル', '件数', '項目数'))
    for s in secs:
        n = len([x for x in s['stmts'] if x['kind'] in ('VALUES', 'SELECT')])
        print('%-30s %-28s %5d %5d' % (s['sheet'],
              ','.join(sorted({x['tbl'] for x in s['stmts'] if x['tbl']})), n, s['ncol']))


if __name__ == '__main__':
    main()
